#!/usr/bin/env python3
"""New filing (2026-10-02): read the receipts dropped on Madder's "New filing" page, and build the patent Filing Report letter
on the firm's letterhead from the REL001 letter (templates/filing-report-letter.docx), keeping its formatting: centered underlined
date, hanging Re:/Subject:/Appl. No.:/Our File No. block, justified body, bold+underlined section headings, underlined dates.
  filing_report.py parse FILE.pdf             -> JSON {kind: ack|payment|assignment|tm_receipt|unknown, ...}
  filing_report.py letter BASE.docx OUT.docx FIELDS.json
FIELDS: date, address_lines[], salutation, title, appl, filed, ref, sections: [{heading: str|null, text: str}, ...]"""
import sys, re, json, copy, os

MONTHS = 'January|February|March|April|May|June|July|August|September|October|November|December'

def pdf_text(path):
    from pypdf import PdfReader
    r = PdfReader(path); out = []
    for pg in r.pages[:4]:
        try: out.append(pg.extract_text(extraction_mode='layout'))
        except Exception: out.append(pg.extract_text() or '')
    return '\n'.join(out)

def iso(mdy):
    m = re.match(r'(\d{2})/(\d{2})/(\d{4})', mdy or ''); return f'{m.group(3)}-{m.group(1)}-{m.group(2)}' if m else ''

def parse(path):
    t = pdf_text(path); flat = re.sub(r'\s+', ' ', t); r = {'kind': 'unknown'}
    if re.search(r'ELECTRONIC ACKNOWLEDGEMENT RECEIPT|ELECTRONIC PAYMENT RECEIPT', t):
        r['kind'] = 'payment' if 'ELECTRONIC PAYMENT RECEIPT' in t else 'ack'
        m = re.search(r'APPLICATION #.*?\n\s*([\d/,]{8,})\s+(\d{2}/\d{2}/\d{4})[^\n]*?([A-Z]{2,5}\d{3}[\w.\-]*)?\s*$', t, re.M | re.S)
        if m: r['app_no'] = re.sub(r'\D', '', m.group(1)); r['date'] = iso(m.group(2)); r['docket'] = (m.group(3) or '').strip()
        m = re.search(r'Title of Invention\s*\n(.*?)\n\s*Application Information', t, re.S)
        if m: r['title'] = re.sub(r'-\s*\n\s*', '-', re.sub(r'\s*\n\s*', ' ', m.group(1).strip())).replace('- ', '-')
        m = re.search(r'APPLICATION TYPE\s+(.*?)\s{2,}PATENT #', flat)
        if m: r['app_type'] = m.group(1).strip()
        m = re.search(r'FIRST NAMED\s+(.+?)\s+INVENTOR', flat)
        if m: r['inventor'] = m.group(1).strip()
        fees = re.findall(r'^\s*(\d{4})\s+(.*?)\s+([\d,]+\.\d{2})\s+\d+\s+([\d,]+\.\d{2})\s*$', t, re.M)
        r['fees'] = [{'code': c, 'desc': d.strip(), 'total': float(tot.replace(',', ''))} for c, d, _, tot in fees]
        r['track_one'] = any(f['code'][1:] == '817' or 'PRIORITIZED' in f['desc'] for f in r['fees'])   # 1817/2817/3817 = Track One request
        codes = [f['code'] for f in r['fees'] if f['code'][0] in '123']
        if codes:   # Patent Center fee codes: 1xxx undiscounted, 2xxx small, 3xxx micro (37 CFR 1.16 schedule)
            p = max(set(c[0] for c in codes), key=lambda d: sum(1 for c in codes if c[0] == d)); r['entity'] = {'1': 'undiscounted', '2': 'small', '3': 'micro'}[p]
        m = re.search(r'TOTAL\s+AMOUNT:\s*\$?([\d,]+\.\d{2})', flat) or re.search(r'TOTAL\s+\$([\d,]+\.\d{2})\s+AMOUNT', flat)
        if m: r['total'] = m.group(1)
    elif 'PATENT ASSIGNMENT COVER SHEET' in t:
        r['kind'] = 'assignment'
        m = re.search(r'CONVEYING PARTY DATA\s+Name\s+Execution Date\s+(.+?)\s+(\d{2}/\d{2}/\d{4})', flat)
        if m: r['conveying'] = m.group(1).strip(); r['executed'] = iso(m.group(2))
        m = re.search(r'RECEIVING PARTY DATA\s+Name:\s+(.+?)\s+Street Address:', flat)
        if m: r['receiving'] = m.group(1).strip()
        g = lambda lab: (re.search(r'RECEIVING PARTY DATA.*?' + lab + r':\s+(.+?)\s+(?:Internal|City|State/Country|Postal Code|PROPERTY NUMBERS)', flat) or [None, ''])[1].strip()
        street, city, st, zp = g('Street Address'), g('City'), g('State/Country'), g('Postal Code')
        country = '' if re.search(r'UNITED STATES', st) else st.title(); state = re.sub(r'\s*UNITED STATES\s*', '', st).title() if re.search(r'UNITED STATES', st) else ''
        r['receiving_address'] = [x for x in [street, ', '.join(x for x in [city, (state + ' ' + zp).strip() if state else zp] if x), country] if x]
        m = re.search(r'Application Number:\s+(\d{8})', flat); r['app_no'] = m.group(1) if m else ''
        m = re.search(r'ATTORNEY DOCKET\s+([A-Z]{2,5}\d{3}[\w.\-]*)\s+NUMBER', flat); r['docket'] = m.group(1) if m else ''
    elif re.search(r'Trademark/service mark application|trademark application', t, re.I) and re.search(r'Serial number:\s*\d{8}', t):
        r['kind'] = 'tm_receipt'
        g = lambda lab: (re.search(r'^\s*' + lab + r':\s*(.+?)\s*$', t, re.M | re.I) or [None, ''])[1]
        r['serial'] = re.sub(r'\D', '', g('Serial number')); r['mark'] = g('Mark'); r['owner'] = g('Owner name'); r['docket'] = g('Docket number')
        fd = re.match(r'(' + MONTHS + r') (\d{1,2}), (\d{4})', g('Filing date'))
        if fd:
            import datetime; r['date'] = datetime.datetime.strptime(fd.group(0), '%B %d, %Y').date().isoformat()
        r['paid'] = g('Amount paid')
    return r

# ---------------------------------------------------------------- the letter
W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
def ptext(e): return ''.join(t.text or '' for t in e.iter(W + 't'))
def put(p, text, start=0):
    rs = p.runs
    if not rs: p.add_run(text); return
    rs[start].text = text
    for r in rs[start + 1:]: r.text = ''

def set_body(p, text):
    """Text into a cloned body paragraph: dates underlined (house style), everything else in the paragraph's first run format."""
    proto = p.runs[0]; underline_proto = proto.underline
    for r in list(p.runs)[1:]: r._element.getparent().remove(r._element)
    parts = re.split(r'((?:' + MONTHS + r') \d{1,2}, \d{4})', text)
    proto.text = parts[0]; proto.underline = underline_proto
    last = proto
    for k, s in enumerate(parts[1:], 1):
        if not s: continue
        e = copy.deepcopy(proto._element); last._element.addnext(e)
        import docx.text.run as dr
        run = dr.Run(e, p); run.text = s; run.underline = bool(k % 2) or underline_proto
        last = run

def letter(base, out, f):
    import docx, docx.text.paragraph as dp
    d = docx.Document(base); P = d.paragraphs
    find = lambda pred: next((p for p in d.paragraphs if pred(p.text)), None)
    p = find(lambda t: re.fullmatch(r'\s*(' + MONTHS + r') \d{1,2}, \d{4}\s*', t))
    if p: put(p, f['date'], next((i for i, r in enumerate(p.runs) if r.text.strip()), 0))
    via = find(lambda t: t.strip().startswith('VIA:'))
    if via is not None:
        ps = []; cur = via._element.getnext()
        while cur is not None and not re.match(r'\s*Re:', ptext(cur)):
            if ptext(cur).strip(): ps.append(cur)
            cur = cur.getnext()
        paras = [dp.Paragraph(e, via._parent) for e in ps]; lines = f['address_lines'] or ['ADDRESS NEEDED']
        gap, inner = paras[-1].paragraph_format.space_after, paras[0].paragraph_format.space_after   # 2026-10-13: the gap above Re: is the last address line's space-after; keep it on whichever line ends up last
        while len(paras) < len(lines):
            e = copy.deepcopy(paras[-2 if len(paras) > 1 else -1]._element); paras[-1]._element.addprevious(e); paras.insert(len(paras) - 1, dp.Paragraph(e, via._parent))
        for para, line in zip(paras, lines): put(para, line)
        for para in paras[len(lines):]: para._element.getparent().remove(para._element)
        for i, para in enumerate(paras[:len(lines)]): para.paragraph_format.space_after = gap if i == len(lines) - 1 else inner
    def field(label, value):
        p = find(lambda t: re.match(r'\s*' + re.escape(label), t))
        if p:
            k = [i for i, r in enumerate(p.runs) if r.text.strip() and not r.text.strip().startswith(label.strip())]
            put(p, value, k[0] if k else len(p.runs) - 1)
            if label.startswith('Appl') and p.runs: set_body_runs_dates(p)
    def set_body_runs_dates(p):
        for r in p.runs:
            if re.search(MONTHS, r.text or '') and not r.underline:
                parts = re.split(r'((?:' + MONTHS + r') \d{1,2}, \d{4})', r.text); r.text = parts[0]; last = r
                import docx.text.run as dr
                for k, s in enumerate(parts[1:], 1):
                    if not s: continue
                    e = copy.deepcopy(r._element); last._element.addnext(e); nr = dr.Run(e, p); nr.text = s; nr.underline = bool(k % 2); last = nr
    title = f['title'] if len(f['title']) <= 88 else f['title'][:86].rsplit(' ', 1)[0] + '...'   # more than two lines in the Re: column: truncated (house style, REL001)
    field('Re:', title); field('Appl. No.:', f['appl'] + ' (Filed: ' + f['filed'] + ')'); field('Our File No.:', f['ref'])
    dear = find(lambda t: t.startswith('Dear ')); put(dear, f['salutation'] or 'Dear Sir or Madam,')
    # body: everything between "Dear ..." and "Sincerely," is rebuilt from the sections
    body, heading = None, None; cur = dear._element.getnext(); old = []
    while cur is not None and 'Sincerely' not in ptext(cur):
        para = dp.Paragraph(cur, dear._parent)
        if ptext(cur).strip():
            if para.runs and para.runs[0].bold and para.runs[0].underline and heading is None: heading = copy.deepcopy(cur)
            elif body is None and para.runs and not para.runs[0].bold: body = copy.deepcopy(cur)
        old.append(cur); cur = cur.getnext()
    spacer = next((e for e in old if not ptext(e).strip()), None); spacer = copy.deepcopy(spacer) if spacer is not None else None
    for e in old: e.getparent().remove(e)
    anchor = dear._element
    def add(proto, text, is_heading):
        nonlocal anchor
        e = copy.deepcopy(proto); anchor.addnext(e); anchor = e; para = dp.Paragraph(e, dear._parent)
        if is_heading: put(para, text)
        else: set_body(para, text)
    if spacer is not None: e = copy.deepcopy(spacer); anchor.addnext(e); anchor = e
    for s in f['sections']:
        if s.get('heading'): add(heading, s['heading'], True)
        add(body, s['text'], False)
    if spacer is not None: e = copy.deepcopy(spacer); anchor.addnext(e)
    d.core_properties.author = os.environ.get('MADDER_ATTORNEY', 'Attorney')
    for r in (r for p in d.paragraphs for r in p.runs if 'Attorney Name, Esq.' in r.text): r.text = r.text.replace('Attorney Name, Esq.', os.environ.get('MADDER_ATTORNEY') or 'Attorney Name, Esq.')   # packaged template's placeholder signer
    d.save(out); dedupe_zip(out)

def dedupe_zip(path):
    """python-docx can write a part twice (docProps/core.xml) when the source .docx already carried a stray copy; Word may then call
    the file damaged. Keep the last copy of each name."""
    import zipfile, os
    with zipfile.ZipFile(path) as z:
        names = [i.filename for i in z.infolist()]
        if len(names) == len(set(names)): return
        last = {i.filename: i for i in z.infolist()}; data = {n: z.read(last[n]) for n in last}
    tmp = path + '.tmp'
    with zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED) as o:
        for n in dict.fromkeys(names): o.writestr(last[n], data[n])
    os.replace(tmp, path)

if __name__ == '__main__':
    if sys.argv[1] == 'parse': print(json.dumps(parse(sys.argv[2])))
    elif sys.argv[1] == 'letter': letter(sys.argv[2], sys.argv[3], json.load(open(sys.argv[4])))
