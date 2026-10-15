#!/usr/bin/env python3
"""Trademark Issuance Letter (2026-10-04: "prepare the issuance letter from the template"), filled into the XYZ004 letter of
2026-08-18 (templates/tm-issuance-letter.docx: Templates/Trademark Issuance Letter.doc wording on the the firm letterhead, as the attorney sent it).
Paragraphs are found by what they say.   tm_issuance_letter.py fill BASE.docx OUT.docx FIELDS.json
FIELDS: date, address_lines[], salutation, mark (Re: line + congratulations), display_text ('' = logo), display_image (path, logo marks),
        reg_no ("8,318,902"), reg_date ("June 23, 2026"), ref, madrid (bool), invoice (bool), reminder ("June 9, 2031"),
        portfolio [[mark, reg_no, reg_date, renewal_opening], ...] (this registration first)."""
import sys, json, copy, io, re, os

def set_text(p, text):
    rs = p.runs
    if not rs: p.add_run(text); return
    rs[0].text = text
    for r in rs[1:]: r._r.getparent().remove(r._r)

def trim(im):
    """The logo without its white border, so the ® sits against the mark itself (2026-10-13)."""
    from PIL import ImageChops, Image
    box = ImageChops.difference(im, Image.new('RGB', im.size, (255, 255, 255))).convert('L').point(lambda v: 255 if v > 24 else 0).getbbox()
    if not box: return im
    m = max(2, round(min(im.size) * 0.01))
    return im.crop((max(0, box[0] - m), max(0, box[1] - m), min(im.width, box[2] + m), min(im.height, box[3] + m)))

def with_mark(im, height=1.25):
    """The logo with the ® drawn at its upper right (2026-10-13, XYZ001: same line as the logo, upper right). Drawn into the
    image because Word and LibreOffice place a raised text run differently. Returns (image, display width in inches); a wide logo is
    scaled to fit the 6.5in line."""
    from PIL import Image, ImageDraw, ImageFont
    import os
    H = 600; im = im.resize((max(1, round(im.width * H / im.height)), H), Image.LANCZOS)
    shown = height * min(1, 5.6 / (height * im.width / H))   # display height once a wide logo is fitted to the line
    font = ImageFont.truetype(os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', 'templates', 'Lexend-SemiBold.ttf'), round(H * 14 / 72 / shown))
    l, t, r, b = font.getbbox('®'); gap = round(H * 0.03)
    out = Image.new('RGB', (im.width + gap + (r - l), H), (255, 255, 255)); out.paste(im, (0, 0))
    ImageDraw.Draw(out).text((im.width + gap - l, -t), '®', font=font, fill=(26, 26, 26))
    return out, min(shown * out.width / H, 6.5)

def set_logo(d, img, sym, im, height=1.25):
    """Puts im (PIL, RGB) in the letter's logo slot with its ®, and drops the separate text ®."""
    import io
    from docx.oxml.ns import qn
    im, w = with_mark(trim(im), height); h = w * im.height / im.width
    part = d.part.related_parts[img._r.xpath('.//a:blip')[0].get(qn('r:embed'))]
    buf = io.BytesIO(); im.save(buf, 'JPEG', quality=95); part._blob = buf.getvalue()
    for e in img._r.xpath('.//wp:extent') + img._r.xpath('.//a:ext'): e.set('cx', str(int(w * 914400))); e.set('cy', str(int(h * 914400)))
    sym._r.getparent().remove(sym._r)

def fill(base, out, f):
    import docx
    from docx.oxml.ns import qn
    from docx.text.paragraph import Paragraph
    d = docx.Document(base); P = d.paragraphs
    find = lambda pre: next(p for p in d.paragraphs if p.text.strip().startswith(pre))
    set_text(P[0], f['date'])
    rec = [p for p in P[2:6]]; lines = f['address_lines'] or ['ADDRESS NEEDED']
    gap, inner = rec[-1].paragraph_format.space_after, rec[0].paragraph_format.space_after   # 2026-10-13: the gap above Re: is the last address line's space-after; keep it on whichever line ends up last
    while len(rec) < len(lines):
        e = copy.deepcopy(rec[-1]._p); rec[-1]._p.addnext(e); rec.append(Paragraph(e, rec[-1]._parent))
    for i, p in enumerate(rec):
        if i < len(lines): set_text(p, lines[i])
        else: p._p.getparent().remove(p._p)
    for i, p in enumerate(rec[:len(lines)]): p.paragraph_format.space_after = gap if i == len(lines) - 1 else inner
    find('Re:').runs[2].text = f['mark']
    find('Reg. No.:').runs[2].text = '%s (%s)' % (f['reg_no'], f['reg_date'])
    find('File No:').runs[2].text = f['ref']
    # 2026-10-13: the File No. value drifted out of line in Word. The label sits at 1620 and the value went to whichever default
    # 720 stop came next, so a narrower label landed it a stop early. An explicit stop at 2880 lines up all four values.
    from docx.shared import Twips
    for pre in ('Re:', 'Subject:', 'Reg. No.:', 'File No:'):
        ts = find(pre).paragraph_format.tab_stops
        if not any(t.position == Twips(2880) for t in ts): ts.add_tab_stop(Twips(2880))
    sal = next(p for p in d.paragraphs if p.text.strip().startswith('Dear '))
    set_text(sal, f['salutation'] or 'Dear Sir or Madam,')
    c = find('Congratulations on the successful'); c.runs[0].text = re.sub('INKA MAR SABOR PERUANO|SAMPLE MARK', lambda _: f['mark'], c.runs[0].text)   # the the firm letter's sample mark / the packaged template's
    disp = next(p for p in d.paragraphs if p._p.xpath('.//w:drawing') and '®' in p.text)
    img, sym = disp.runs[0], disp.runs[1]
    if f.get('display_text') or not f.get('display_image'):
        img._r.getparent().remove(img._r); sym.text = (f.get('display_text') or f['mark']) + '®'
    else:
        from PIL import Image
        set_logo(d, img, sym, Image.open(f['display_image']).convert('RGB'))
    r = find('Finally, for a trademark registration')
    t = re.sub('June 9, 2031|January 1, 2031', lambda _: f['reminder'], r.runs[0].text).replace('we have scheduled to send you', 'we have docketed to send you')   # no-reminder-promises (2026-10-04)
    if f.get('madrid'): t = t.replace('(“Section 8 Affidavit”)', '(“Section 71 Affidavit”)')
    r.runs[0].text = t
    if f.get('invoice'):
        a = find('Again, congratulations.'); a.runs[0].text = a.runs[0].text.rstrip() + '  My invoice is enclosed.'
    tbl = d.tables[0]; tmpl = tbl.rows[1]._tr; old = [row._tr for row in tbl.rows[1:]]
    for vals in f['portfolio']:
        tr = copy.deepcopy(tmpl); tbl._tbl.append(tr); row = tbl.rows[-1]
        trPr = tr.get_or_add_trPr(); trPr.append(trPr.makeelement(qn('w:cantSplit'), {}))
        for cell, v in zip(row.cells, vals[:4]): set_text(cell.paragraphs[0], v)
        if len(vals) > 4 and vals[4]:   # maintenance window open / opening within 3 months / in grace (2026-10-04): highlighted row
            for cell in row.cells:
                tcPr = cell._tc.get_or_add_tcPr()
                for old_shd in tcPr.findall(qn('w:shd')): tcPr.remove(old_shd)
                tcPr.append(tcPr.makeelement(qn('w:shd'), {qn('w:val'): 'clear', qn('w:color'): 'auto', qn('w:fill'): 'FFF2CC'}))
    for tr in old: tbl._tbl.remove(tr)
    if f.get('portfolio_note'):   # "Please note the highlighted row(s) above. ..." before the closing paragraph
        a = find('Again, congratulations.'); np_ = copy.deepcopy(a._p); a._p.addprevious(np_)
        from docx.text.paragraph import Paragraph
        set_text(Paragraph(np_, a._parent), f['portfolio_note'])
    for r in (r for p in d.paragraphs for r in p.runs if 'Attorney Name, Esq.' in r.text): r.text = r.text.replace('Attorney Name, Esq.', os.environ.get('MADDER_ATTORNEY') or 'Attorney Name, Esq.')   # packaged template's placeholder signer
    d.save(out)

if __name__ == '__main__':
    if len(sys.argv) == 5 and sys.argv[1] == 'fill': fill(sys.argv[2], sys.argv[3], json.load(open(sys.argv[4]))); print('ok')
    else: sys.exit(__doc__)
