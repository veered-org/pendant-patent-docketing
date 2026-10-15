#!/usr/bin/env python3
"""Patent Issuance Letter (2026-10-02), filled into the firm's XYZ005 letter (templates/patent-issuance-letter.docx) so the
letterhead, fonts and underlines stay exactly as Word saved them. Paragraphs are found by what they say, not by position.
  issuance_letter.py fill BASE.docx OUT.docx FIELDS.json
  issuance_letter.py address LETTER.(docx|pdf)     -> JSON {lines: [...], salutation: "Dear Mr. X,"} or {} when not found
FIELDS: date, address_lines[], salutation, title, appl, filed, patent_no, issued, ref, design (bool), pta_days (int),
        expires ("Wednesday, March 1, 2045"), reminder (first-fee window opens, "January 1, 2029"), claims_ok (bool), filename,
        priority ("June 20, 2017": earliest non-provisional priority date when earlier than the filing date, else "")"""
import sys, json, re, copy, os

CLAIM_SENTENCE = re.compile(r'\s*The independent claim language in your Patent has been checked for accuracy and has been found to be correct\.')

def put(p, text, start=0):
    """Text into run START (its formatting kept); later runs emptied."""
    rs = p.runs
    if not rs: p.add_run(text); return
    rs[start].text = text
    for r in rs[start + 1:]: r.text = ''

def first_text_run(p):
    for i, r in enumerate(p.runs):
        if r.text.strip(): return i
    return 0

def remove(p): p._element.getparent().remove(p._element)

def fill(base, out, f):
    import docx
    d = docx.Document(base); P = d.paragraphs
    find = lambda pred: next((p for p in d.paragraphs if pred(p.text)), None)
    # date (right-aligned by tabs)
    p = find(lambda t: re.fullmatch(r'\t*[A-Z][a-z]+ \d{1,2}, \d{4}\s*', t))
    if p: put(p, f['date'], first_text_run(p))
    # address block: between "VIA: E-MAIL ONLY" and "Re:"
    via = find(lambda t: t.strip().startswith('VIA:'))
    if via is not None:
        ps = []; cur = via._element.getnext()
        while cur is not None:
            txt = ''.join(t.text or '' for t in cur.iter('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t'))
            if txt.strip().startswith('Re:') or re.match(r'\s*Re:', txt): break
            if txt.strip(): ps.append(cur)
            cur = cur.getnext()
        import docx.text.paragraph as dp
        paras = [dp.Paragraph(e, via._parent) for e in ps]
        lines = f['address_lines'] or ['ADDRESS NEEDED']
        gap, inner = paras[-1].paragraph_format.space_after, paras[0].paragraph_format.space_after   # 2026-10-13: the gap above Re: is the last address line's space-after; keep it on whichever line ends up last
        while len(paras) < len(lines):
            e = copy.deepcopy(paras[-1]._element); paras[-1]._element.addnext(e); paras.append(dp.Paragraph(e, via._parent))
        for para, line in zip(paras, lines): put(para, line)
        for para in paras[len(lines):]: remove(para)
        for i, para in enumerate(paras[:len(lines)]): para.paragraph_format.space_after = gap if i == len(lines) - 1 else inner
    # Re: title (the sample wraps it onto a second paragraph: that one goes)
    re_p = find(lambda t: re.match(r'\s*Re:', t))
    if re_p is not None:
        k = [i for i, r in enumerate(re_p.runs) if r.text.strip() and r.text.strip() != 'Re:']
        put(re_p, f['title'], k[0] if k else len(re_p.runs) - 1)
        nxt = re_p._element.getnext()
        import docx.text.paragraph as dp
        if nxt is not None:
            np_ = dp.Paragraph(nxt, re_p._parent)
            if np_.text.strip() and 'Subject:' not in np_.text: remove(np_)
    p = find(lambda t: 'Appl. No.:' in t)
    if p: k = [i for i, r in enumerate(p.runs) if re.search(r'\d', r.text)]; put(p, f['appl'] + ' (Filed: ' + f['filed'] + ')', k[0] if k else len(p.runs) - 1)
    p = find(lambda t: 'Patent No.:' in t)
    if p: k = [i for i, r in enumerate(p.runs) if re.search(r'\d', r.text)]; put(p, '\t' + f['patent_no'] + ' (Issued: ' + f['issued'] + ')', k[0] if k else len(p.runs) - 1)
    p = find(lambda t: re.search(r'File:', t))
    if p: k = [i for i, r in enumerate(p.runs) if re.search(r'[A-Z]{3}\d', r.text)]; put(p, '\t' + f['ref'], k[0] if k else len(p.runs) - 1)
    p = find(lambda t: t.startswith('Dear '))
    if p: put(p, f['salutation'] or 'Dear Sir or Madam,')
    p = find(lambda t: t.startswith('Congratulations!'))
    if p and not f.get('claims_ok'):
        for r in p.runs: r.text = CLAIM_SENTENCE.sub('', r.text)
    p = find(lambda t: t.startswith('In order to collect'))
    if p:
        for r in p.runs:
            if re.fullmatch(r'D?\d{1,2},\d{3},\d{3}', r.text.strip()): r.text = f['patent_no']
        if f['patent_no'] not in p.text: put(p, re.sub(r'U\.S\. Patent [D\d,]+', 'U.S. Patent ' + f['patent_no'], p.text))
    term = find(lambda t: t.startswith('If all maintenance fees'))
    mf = find(lambda t: t.startswith('Referring again to the maintenance fees'))
    if term:
        rs = term.runs; u = [i for i, r in enumerate(rs) if r.underline and r.text.strip()]
        if f.get('design'):
            put(term, 'Your Patent will remain in force until ', 0)
            if u:
                rs = term.runs; rs[u[0]].text = f['expires']
                plain = [r for r in rs[u[0] + 1:] if not r.underline]
                for r in rs[u[0] + 1:]: r.text = ''
                if plain: plain[-1].text = '.'      # the period goes in a run that is not underlined
            else: put(term, 'Your Patent will remain in force until ' + f['expires'] + '.')
        else:
            for r in rs:
                if re.fullmatch(r'\s*and \d+ days\s*', r.text): r.text = (' and {:,} day{}'.format(f['pta_days'], '' if f['pta_days'] == 1 else 's')) if f['pta_days'] else ''
                if f.get('priority') and 'from the date of filing' in r.text: r.text = r.text.replace('from the date of filing', 'from the earliest priority filing date (%s)' % f['priority'])
            if u: rs[u[0]].text = f['expires']; [setattr(rs[i], 'text', '') for i in u[1:]]
    if mf:
        if f.get('design'):
            sp = mf._element.getnext(); remove(mf)
            if sp is not None and not ''.join(t.text or '' for t in sp.iter('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t')).strip(): sp.getparent().remove(sp)
        else:
            rs = mf.runs; u = [i for i, r in enumerate(rs) if r.underline and re.search(r'\d|[A-Z][a-z]+ ', r.text) and 'date of issue' not in r.text]
            if u: rs[u[0]].text = f['reminder']; [setattr(rs[i], 'text', '') for i in u[1:]]
            # 2026-10-05: the date is when the first fee's window opens (3 years from issue), not a reminder shortly before 3.5 years
            for r in rs:
                r.text = r.text.replace('Thus, a reminder is scheduled to be sent to you regarding the first maintenance fee shortly before 3.5 yrs from the issue date of your Patent, or on about ',
                                        'Thus, the window to pay the first maintenance fee opens 3 yrs from the issue date of your Patent, on ')
                if r.text == '.  ' and 'window to pay the first' in mf.text and 'We have docketed' not in mf.text: r.text = '.  We have docketed to send you a reminder at that time.  '
    p = find(lambda t: re.search(r'-\d*[a-z]?-?Issuance Report\.doc|Issuance Report\.doc', t))
    if p: put(p, f['filename'], first_text_run(p))
    d.core_properties.author = os.environ.get('PENDANT_ATTORNEY', 'Attorney')
    for r in (r for p in d.paragraphs for r in p.runs if 'Attorney Name, Esq.' in r.text): r.text = r.text.replace('Attorney Name, Esq.', os.environ.get('PENDANT_ATTORNEY') or 'Attorney Name, Esq.')   # packaged template's placeholder signer
    d.save(out); dedupe_zip(out)

def address(path):
    if path.lower().endswith('.docx'):
        import docx; lines = [p.text for p in docx.Document(path).paragraphs]
    else:
        from pypdf import PdfReader; lines = (PdfReader(path).pages[0].extract_text() or '').split('\n')
    lines = [l.replace('\xa0', ' ').strip() for l in lines]
    i = next((k for k, l in enumerate(lines) if re.match(r'VIA:\s*E-?MAIL', l, re.I)), None)
    if i is None: return {}
    block = []
    for l in lines[i + 1:]:
        if re.match(r'Re:', l) or re.match(r'\s*Re\s*:', l): break
        if l: block.append(l)
        if len(block) > 7: return {}
    sal = next((l for l in lines if re.match(r'Dear .{2,60},$', l)), '')
    return {'lines': block, 'salutation': sal} if block else {}

def dedupe_zip(path):
    """python-docx can write a part twice (docProps/core.xml) when the source .docx already carried a stray copy; Word may then call
    the file damaged. Keep the last copy of each name."""
    import zipfile, os
    with zipfile.ZipFile(path) as z:
        names = [i.filename for i in z.infolist()]
        if len(names) == len(set(names)): return
        last = {i.filename: i for i in z.infolist()}; data = {n: z.read(last[n]) for n in last}
    tmp = path + '.tmp'
    with zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED) as o:
        for n in dict.fromkeys(names): o.writestr(last[n], data[n])
    os.replace(tmp, path)

if __name__ == '__main__':
    if sys.argv[1] == 'fill': fill(sys.argv[2], sys.argv[3], json.load(open(sys.argv[4])))
    elif sys.argv[1] == 'address': print(json.dumps(address(sys.argv[2])))
